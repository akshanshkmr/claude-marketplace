#!/usr/bin/env python3
"""
mechanical_edits.py — JMIR deterministic copyediting pass.

WHY THIS EXISTS
---------------
A human (or an LLM) reading 6 reference files will inevitably miss a few of the
~100 routine find/replace edits JMIR copyeditors make on every paper. A regex
pass does not get tired and does not forget. So the division of labour is:

    * This script   -> the DETERMINISTIC layer. Every mechanical edit that can be
                       expressed as a pattern is caught here, exhaustively.
    * The skill +   -> the JUDGEMENT layer. Rephrasing, restructuring, deciding
      the agent        whether a context-dependent swap fits, statistical
                       completeness, and author queries.

Findings are split into two confidence levels:
    auto   = safe to apply verbatim (e.g. "utilization" -> "use", P=0.03 -> P=.03)
    review = correct most of the time but context-dependent; PROPOSE it and let a
             human/agent confirm (e.g. "subjects" -> "participants", which is wrong
             inside "subject matter"). Many `review` items also map to an author
             query in query-bank.md.

USAGE
-----
    python mechanical_edits.py MANUSCRIPT.docx                 # human-readable report
    python mechanical_edits.py MANUSCRIPT.docx --json OUT.json # structured findings
    python mechanical_edits.py MANUSCRIPT.docx --apply OUT.docx# write a copy with AUTO fixes applied
    python mechanical_edits.py MANUSCRIPT.txt                  # plain-text also accepted

The JSON output is designed to be consumed by an agent that then (a) applies the
edits as tracked changes in its own harness and (b) layers on judgement edits +
queries. The --apply mode applies only `auto` fixes as plain text (use Word's
Review > Compare against the original to regenerate tracked changes if needed).

The rule tables below are DATA, not logic — extend them freely as house style
evolves. Each rule carries a short `note` that becomes the edit's rationale.
"""

from __future__ import annotations
import argparse
import json
import re
import sys
from dataclasses import dataclass, asdict
from pathlib import Path

# ──────────────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────────────

def _match_case(template: str, replacement: str) -> str:
    """Return `replacement` cased to mirror `template` (Title/UPPER/lower)."""
    if template.isupper() and len(template) > 1:
        return replacement.upper()
    if template[:1].isupper():
        return replacement[:1].upper() + replacement[1:]
    return replacement


@dataclass
class Finding:
    rule_id: str
    category: str
    mode: str          # "auto" | "review"
    original: str
    suggestion: str
    note: str
    section: str
    para_index: int
    context: str

# ──────────────────────────────────────────────────────────────────────────────
# RULE TABLES
# Each entry: (rule_id, regex, replacement, note)
# `replacement` may contain backrefs (\1) or the sentinel {KEEPCASE} handled below.
# ──────────────────────────────────────────────────────────────────────────────

# British -> American spelling (conservative, high-confidence subset). AUTO.
# Ambiguous noun/verb pairs (licence/license, practise/practice, programme) are
# deliberately placed in REVIEW further down.
_SPELL = {
    r"colour": "color", r"colours": "colors", r"coloured": "colored",
    r"behaviour": "behavior", r"behaviours": "behaviors", r"behavioural": "behavioral",
    r"favour": "favor", r"favours": "favors", r"favourite": "favorite",
    r"labour": "labor",
    r"odour": "odor", r"tumour": "tumor", r"tumours": "tumors",
    r"organise": "organize", r"organised": "organized", r"organising": "organizing",
    r"organisation": "organization", r"organisations": "organizations",
    r"analyse": "analyze", r"analysed": "analyzed", r"analysing": "analyzing",
    r"catalogue": "catalog", r"dialogue": "dialogue",  # dialogue kept (AMA allows)
    r"centre": "center", r"centres": "centers", r"centred": "centered",
    r"fibre": "fiber", r"litre": "liter", r"litres": "liters",
    r"metre": "meter", r"metres": "meters",
    r"defence": "defense", r"offence": "offense",
    r"grey": "gray", r"greyscale": "grayscale",
    r"mould": "mold", r"sulphur": "sulfur",
    r"oedema": "edema", r"foetal": "fetal", r"foetus": "fetus",
    r"haemoglobin": "hemoglobin", r"haematology": "hematology",
    r"paediatric": "pediatric", r"paediatrics": "pediatrics",
    r"anaemia": "anemia", r"anaemic": "anemic",
    r"ageing": "aging",
    r"modelling": "modeling", r"modelled": "modeled",
    r"labelling": "labeling", r"labelled": "labeled",
    r"travelling": "traveling", r"travelled": "traveled",
    r"fulfil": "fulfill", r"fulfilment": "fulfillment",
    r"enrol": "enroll", r"enrolment": "enrollment",
    r"towards": "toward", r"amongst": "among", r"whilst": "while",
    r"learnt": "learned", r"summarise": "summarize", r"summarised": "summarized",
    r"recognise": "recognize", r"recognised": "recognized",
    r"prioritise": "prioritize", r"prioritised": "prioritized",
    r"utilise": "use", r"utilised": "used", r"utilising": "using",  # also a word-swap
}
SPELLING_RULES = [
    (f"spell.{i}", re.compile(rf"\b{pat}\b", re.IGNORECASE), rep,
     "American (US) spelling per JMIR house style")
    for i, (pat, rep) in enumerate(_SPELL.items())
]

# AUTO word-swaps — safe regardless of context. (rule_id, pattern, repl, note)
WORD_SWAPS_AUTO = [
    ("swap.utiliz", re.compile(r"\butiliz(e|es|ed|ing|ation)\b", re.IGNORECASE),
     lambda m: _match_case(m.group(0), {"e": "use", "es": "uses", "ed": "used",
                                        "ing": "using", "ation": "use"}[m.group(1).lower()]),
     'JMIR avoids "utilize/utilization"; use "use"'),
    ("swap.healthcare", re.compile(r"\bhealthcare\b", re.IGNORECASE), "health care",
     'JMIR style: "health care" (two words)'),
    ("swap.homecare", re.compile(r"\bhomecare\b", re.IGNORECASE), "home care",
     'JMIR style: "home care" (two words)'),
    ("swap.sociodemo", re.compile(r"\bsocio[-\s]demographic\b", re.IGNORECASE), "sociodemographic",
     'JMIR style: "sociodemographic" (one word, no hyphen)'),
    ("swap.wellbeing", re.compile(r"\bwell\s?being\b", re.IGNORECASE), "well-being",
     'JMIR style: "well-being" (hyphenated)'),
    ("swap.firstly", re.compile(r"\bfirstly\b", re.IGNORECASE), "first",
     'Use "first" not "firstly"'),
    ("swap.secondly", re.compile(r"\bsecondly\b", re.IGNORECASE), "second",
     'Use "second" not "secondly"'),
    ("swap.thirdly", re.compile(r"\bthirdly\b", re.IGNORECASE), "third",
     'Use "third" not "thirdly"'),
    ("swap.lastly", re.compile(r"\blastly\b", re.IGNORECASE), "finally",
     'Use "finally" not "lastly"'),
    ("swap.priorto", re.compile(r"\bprior to\b", re.IGNORECASE), "before",
     'Use "before" not "prior to"'),
    ("swap.ehealth", re.compile(r"\be[-\s]health\b", re.IGNORECASE), "eHealth",
     'JMIR neologism style: "eHealth" (not e-health/e-Health)'),
    ("swap.mhealth", re.compile(r"\bm[-\s]health\b", re.IGNORECASE), "mHealth",
     'JMIR neologism style: "mHealth" (not m-health/m-Health)'),
    ("swap.esource", re.compile(r"\be[-\s]source\b", re.IGNORECASE), "eSource",
     'JMIR neologism style: "eSource"'),
    ("swap.cronbach", re.compile(r"\bCronbach'?s?\s+(alpha|α)\b"), r"Cronbach \1",
     "AMA: no possessive on statistical eponyms (Cronbach α)"),
    ("swap.cohen", re.compile(r"\bCohen'?s\s+d\b"), "Cohen d",
     "AMA: no possessive on statistical eponyms (Cohen d)"),
    ("swap.fisher", re.compile(r"\bFisher'?s\s+exact\b"), "Fisher exact",
     "AMA: no possessive on statistical eponyms (Fisher exact)"),
    ("swap.hedges", re.compile(r"\bHedges?'?s?\s+g\b"), "Hedges g",
     "AMA: no possessive on statistical eponyms (Hedges g)"),
    ("time.noon", re.compile(r"\b12(:00)?\s*[Pp]\.?\s?[Mm]\.?"), "noon",
     'JMIR: write "noon" for 12 PM'),
    ("time.midnight", re.compile(r"\b12(:00)?\s*[Aa]\.?\s?[Mm]\.?"), "midnight",
     'JMIR: write "midnight" for 12 AM'),
    ("tm.symbols", re.compile(r"[®™℠]"), "",
     "Remove trademark/registered/service-mark symbols"),
]

# REVIEW word-swaps — usually right, but confirm in context. Many pair with a query.
WORD_SWAPS_REVIEW = [
    ("rev.employ", re.compile(r"\bemploy(s|ed|ing)?\b", re.IGNORECASE), "use",
     'JMIR avoids "employ" in the sense of "use" (keep if it means to hire)'),
    ("rev.suffer", re.compile(r"\bsuffer(s|ed|ing)?\b(\s+from)?", re.IGNORECASE), "experience",
     'AMA: avoid "suffer from"; prefer "experience" / "have" (people-first)'),
    ("rev.subjects", re.compile(r"\bsubjects\b", re.IGNORECASE), "participants",
     'JMIR prefers "participants" over "subjects" (not when it means subject matter)'),
    ("rev.victim", re.compile(r"\bvictims?\b", re.IGNORECASE), "survivor",
     'AMA: prefer "survivor" for those who survive violence/disaster'),
    ("rev.application", re.compile(r"\bapplications?\b", re.IGNORECASE), "app",
     'JMIR shortens "application" to "app" when it means software (not a job/method application)'),
    ("rev.elderly", re.compile(r"\belderly\b", re.IGNORECASE), "older people",
     'Avoid "elderly"; use "older people" / "older adults" (query the author)'),
    ("rev.comparedto", re.compile(r"\bcompared to\b", re.IGNORECASE), "compared with",
     'AMA: "compared with" for differences (keep "compared to" only for likening)'),
    ("rev.further", re.compile(r"\bFurther\b,", ), "Furthermore,",
     'As a sentence-opening connector, prefer "Furthermore," over "Further,"'),
    ("rev.additionally", re.compile(r"\bAdditionally\b,"), "In addition,",
     'Prefer "In addition," over "Additionally,"'),
    ("rev.moreover", re.compile(r"\bMoreover\b,"), "Furthermore,",
     'Prefer "Furthermore," / "In addition," over "Moreover,"'),
    ("rev.asaresult", re.compile(r"\bas a result\b", re.IGNORECASE), "thus",
     'Prefer "thus" over "as a result"'),
    ("rev.though", re.compile(r",\s+though\b", re.IGNORECASE), ", however",
     'Prefer "however" over trailing "though"'),
    ("rev.presentstudy", re.compile(r"\b(the\s+)?(present|current)\s+(study|paper|work|article)\b", re.IGNORECASE),
     "this study", 'JMIR: "this study"/"this paper", not "present/current study"'),
    ("rev.manuscript", re.compile(r"\bmanuscripts?\b", re.IGNORECASE), "paper",
     'JMIR prefers "paper"/"study" over "manuscript" (KEEP in Acknowledgments)'),
    ("rev.article", re.compile(r"\barticles?\b", re.IGNORECASE), "paper",
     'JMIR prefers "paper"/"study" over "article" (when referring to this work)'),
    ("rev.modtosevere", re.compile(r"\bmoderate to severe\b", re.IGNORECASE), "vigorous",
     'In exercise-intensity context JMIR prefers "vigorous" (confirm context)'),
    ("rev.relatives", re.compile(r"\brelatives\b", re.IGNORECASE), "caregivers",
     'Consider "caregivers" if that is the intended meaning (query)'),
    ("rev.medprof", re.compile(r"\bmedical professionals?\b", re.IGNORECASE), "health care professional",
     'Prefer "health care professional" for consistency'),
    ("rev.healthprof", re.compile(r"\bhealth professionals?\b", re.IGNORECASE), "health care professionals",
     'Prefer "health care professionals"'),
    ("rev.peoplefirst", re.compile(r"\b([A-Za-z]+)-infected (patients|people|individuals)\b"),
     r"\2 infected with \1",
     "People-first language: e.g. 'HIV-infected patients' -> 'patients infected with HIV'"),
    # noun/verb-ambiguous British spellings -> review, not auto
    ("rev.licence", re.compile(r"\blicence\b", re.IGNORECASE), "license",
     'US spelling "license" (confirm noun vs verb)'),
    ("rev.practise", re.compile(r"\bpractise\b", re.IGNORECASE), "practice",
     'US spelling "practice" (confirm noun vs verb)'),
    ("rev.programme", re.compile(r"\bprogramme\b", re.IGNORECASE), "program",
     'US spelling "program"'),
]

# Statistics normalisation. AUTO unless noted.
STATS_AUTO = [
    # Leading-zero removal + no spaces around operator for P values.
    ("stat.pvalue", re.compile(r"\bP\s*([=<>])\s*0?\.(\d+)"),
     lambda m: f"P{m.group(1)}.{m.group(2)}",
     "P values: no leading zero, no spaces around operator (P=.03, P<.001)"),
    ("stat.pvalue_lc", re.compile(r"\bp\s*([=<>])\s*0?\.(\d+)"),
     lambda m: f"P{m.group(1)}.{m.group(2)}",
     'Capitalize and format P value (italic P, no leading zero)'),
    # alpha/beta leading zero
    ("stat.alpha", re.compile(r"\b(α|β)\s*=\s*0?\.(\d+)"),
     lambda m: f"{m.group(1)}=.{m.group(2)}",
     "α/β values: no leading zero, no spaces around operator"),
]
STATS_REVIEW = [
    ("stat.ptoolow", re.compile(r"\bP\s*[=<]\s*0?\.000\d*"), "P<.001",
     "P values below .001 are written as P<.001"),
    ("stat.pequals1", re.compile(r"\bP\s*=\s*1(\.0+)?\b"), "P≥.99",
     'P=1 should be written P≥.99'),
    ("stat.trend", re.compile(r"\btrend(ed|ing)?\s+toward(s)?\s+significance\b", re.IGNORECASE),
     "[state the trend, then whether it was statistically significant + the value]",
     'Avoid "trending toward significance"; state the result and its significance (query)'),
    ("stat.approached", re.compile(r"\bapproached significance\b", re.IGNORECASE),
     "[state significance explicitly]", 'Avoid "approached significance" (query)'),
    ("stat.kgm2", re.compile(r"\bkg[.\u00b7]m-?2\b"), "kg/m²",
     'Unit format: kg/m² (solidus, superscript 2)'),
    ("stat.nyearperiod", re.compile(r"\b(\d+)\s+(year|month|week|day)-(period|old|interval)\b", re.IGNORECASE),
     r"\1-\2 \3", 'Hyphenate the modifier: e.g. "3 year-period" -> "3-year period"'),
]

# Phrases & punctuation to flag for review/query.
FLAG_REVIEW = [
    ("flag.andor", re.compile(r"\band/or\b", re.IGNORECASE), '"and" or "or"',
     'JMIR avoids "and/or"; choose "and" or "or"'),
    ("flag.url", re.compile(r"https?://\S+"), "[move to RefCheck or remove]",
     "URLs are not allowed in the running text/abstract; convert to a reference or remove"),
    ("flag.sms", re.compile(r"\bSMS\b(?!\s+text)"), "SMS text messaging",
     '"SMS" should not stand alone: "SMS text messaging" (abstract) / "text messaging" (title)'),
    ("flag.banned_ongoing", re.compile(r"\b(ongoing|currently underway|to date|at the time of (this )?(writing|publication))\b", re.IGNORECASE),
     "[time-bound phrasing — rephrase for a long-term readership]",
     'Avoid "ongoing/currently underway/to date/at the time of writing"'),
    ("flag.normal", re.compile(r"\b(ab)?normal\b", re.IGNORECASE), "[describe specifically]",
     'AMA: avoid "normal"/"abnormal" for a person\'s health status (query)'),
    ("flag.currency", re.compile(r"(?<![A-Za-z])\$\s?\d"), "US $",
     'Specify currency: "$" should be "US $" or give the conversion rate (query)'),
    # Note: a generic "solidus" detector flags too many legitimate slashes
    # (race/ethnicity, dates, units, URLs). The high-value case — "/" meaning
    # "and"/"or" — is covered by flag.andor above; leave the rest to human review.
]

# Hyphenation: suggest hyphenating these compounds when used attributively. REVIEW.
HYPHEN_COMPOUNDS = [
    ("quality of life", "quality-of-life"),
    ("standard of care", "standard-of-care"),
    ("intention to treat", "intention-to-treat"),
    ("proof of concept", "proof-of-concept"),
    ("one size fits all", "one-size-fits-all"),
    ("face to face", "face-to-face"),
    ("step by step", "step-by-step"),
    ("day to day", "day-to-day"),
    ("around the clock", "around-the-clock"),
    ("gain of function", "gain-of-function"),
    ("loss of function", "loss-of-function"),
]
HYPHEN_RULES = [
    (f"hyph.{i}", re.compile(rf"\b{re.escape(a)}\b", re.IGNORECASE), b,
     f'Hyphenate "{a}" -> "{b}" when used before a noun (review)')
    for i, (a, b) in enumerate(HYPHEN_COMPOUNDS)
]

# Section names used to scope a couple of context-sensitive rules.
_ACK_RE = re.compile(r"^\s*acknowledg", re.IGNORECASE)
_ABSTRACT_RE = re.compile(r"^\s*abstract\b", re.IGNORECASE)
_HEADING_HINT = re.compile(r"^\s*(abstract|introduction|background|objective|methods?|results?|"
                           r"discussion|conclusions?|acknowledgments?|references?|"
                           r"conflicts of interest|funding|authors[\u2019']? contributions|"
                           r"data availability|abbreviations|ethical considerations)\b", re.IGNORECASE)

ALL_RULE_GROUPS = [
    ("spelling", SPELLING_RULES, "auto"),
    ("word-swap", WORD_SWAPS_AUTO, "auto"),
    ("word-swap", WORD_SWAPS_REVIEW, "review"),
    ("statistics", STATS_AUTO, "auto"),
    ("statistics", STATS_REVIEW, "review"),
    ("phrase-flag", FLAG_REVIEW, "review"),
    ("hyphenation", HYPHEN_RULES, "review"),
]

# ──────────────────────────────────────────────────────────────────────────────
# Core scanning
# ──────────────────────────────────────────────────────────────────────────────

def _read_paragraphs(path: Path) -> list[str]:
    if path.suffix.lower() in {".docx"}:
        import docx
        doc = docx.Document(str(path))
        paras = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    paras.append(cell.text)
        return paras
    return path.read_text(encoding="utf-8", errors="ignore").splitlines()


def _apply_repl(repl, m) -> str:
    """repl may be a str (with optional backrefs) or a callable(match)->str."""
    if callable(repl):
        return repl(m)
    # expand backrefs, then mirror case for single-word literal replacements
    out = m.expand(repl) if "\\" in repl else repl
    if "\\" not in repl and m.group(0).isalpha() and out.replace("-", "").isalpha():
        out = _match_case(m.group(0), out)
    return out


def scan(paragraphs: list[str]) -> list[Finding]:
    findings: list[Finding] = []
    section = "(start)"
    for i, para in enumerate(paragraphs):
        text = para
        if not text.strip():
            continue
        if _HEADING_HINT.match(text.strip()) and len(text.strip()) < 60:
            section = text.strip()
        url_spans = [mu.span() for mu in re.finditer(r"https?://\S+", text)]
        for category, group, mode in ALL_RULE_GROUPS:
            for rule_id, pattern, repl, note in group:
                for m in pattern.finditer(text):
                    orig = m.group(0)
                    # don't fire spelling/swap/flag rules on text inside a URL
                    if rule_id != "flag.url" and any(s <= m.start() < e for s, e in url_spans):
                        continue
                    # "normal" is fine in statistical usage (normal distribution,
                    # normality, normalize, normally, lognormal) — only flag it as a
                    # description of a person's health status.
                    if rule_id == "flag.normal":
                        tail = text[m.end():m.end() + 14].lower()
                        prev = text[max(0, m.start() - 3):m.start()].lower()
                        if tail.startswith(("ity", "ize", "iza", "ly", " distrib")) or prev == "log":
                            continue
                    # context-sensitive suppression
                    if rule_id == "rev.manuscript" and _ACK_RE.search(section):
                        continue  # "manuscript" is allowed in Acknowledgments
                    suggestion = _apply_repl(repl, m)
                    if suggestion == orig:
                        continue
                    start = max(0, m.start() - 35)
                    end = min(len(text), m.end() + 35)
                    ctx = ("…" if start else "") + text[start:end] + ("…" if end < len(text) else "")
                    findings.append(Finding(
                        rule_id=rule_id, category=category, mode=mode,
                        original=orig, suggestion=suggestion, note=note,
                        section=section, para_index=i, context=ctx.replace("\n", " "),
                    ))
    return findings


def apply_auto(paragraphs: list[str]) -> tuple[list[str], int]:
    """Apply only AUTO rules to plain text. Returns (new_paragraphs, n_changes)."""
    n = 0
    out = []
    section = "(start)"
    for para in paragraphs:
        text = para
        if _HEADING_HINT.match(text.strip()) and len(text.strip()) < 60:
            section = text.strip()
        for category, group, mode in ALL_RULE_GROUPS:
            if mode != "auto":
                continue
            for rule_id, pattern, repl, note in group:
                def _sub(m):
                    nonlocal n
                    rep = _apply_repl(repl, m)
                    if rep != m.group(0):
                        n += 1
                    return rep
                text = pattern.sub(_sub, text)
        out.append(text)
    return out, n


def _write_docx_auto(src: Path, dst: Path) -> int:
    import docx
    doc = docx.Document(str(src))
    n = 0
    section = "(start)"

    def fix_run_text(t: str) -> str:
        nonlocal n
        for category, group, mode in ALL_RULE_GROUPS:
            if mode != "auto":
                continue
            for rule_id, pattern, repl, note in group:
                def _sub(m):
                    nonlocal n
                    rep = _apply_repl(repl, m)
                    if rep != m.group(0):
                        n += 1
                    return rep
                t = pattern.sub(_sub, t)
        return t

    for p in doc.paragraphs:
        for run in p.runs:
            if run.text:
                run.text = fix_run_text(run.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text:
                            run.text = fix_run_text(run.text)
    doc.save(str(dst))
    return n


# ──────────────────────────────────────────────────────────────────────────────
# Reporting
# ──────────────────────────────────────────────────────────────────────────────

def report(findings: list[Finding]) -> str:
    auto = [f for f in findings if f.mode == "auto"]
    review = [f for f in findings if f.mode == "review"]
    lines = []
    lines.append("=" * 78)
    lines.append("JMIR MECHANICAL-EDITS REPORT")
    lines.append("=" * 78)
    lines.append(f"  {len(auto)} auto-fixable + {len(review)} review/query = {len(findings)} total findings")
    lines.append("")
    for label, group in (("AUTO (safe to apply)", auto), ("REVIEW / QUERY (confirm in context)", review)):
        lines.append("-" * 78)
        lines.append(f"{label}: {len(group)}")
        lines.append("-" * 78)
        by_cat: dict[str, list[Finding]] = {}
        for f in group:
            by_cat.setdefault(f"{f.category} · {f.note}", []).append(f)
        for key in sorted(by_cat):
            items = by_cat[key]
            lines.append(f"\n• {key}  ({len(items)}×)")
            for f in items[:8]:
                lines.append(f"    p{f.para_index} [{f.section[:24]}]  "
                             f"\"{f.original}\" → \"{f.suggestion}\"")
                lines.append(f"        … {f.context}")
            if len(items) > 8:
                lines.append(f"    … and {len(items) - 8} more")
        lines.append("")
    return "\n".join(lines)


def main(argv=None):
    ap = argparse.ArgumentParser(description="JMIR deterministic mechanical-edits pass")
    ap.add_argument("manuscript", help="path to .docx or .txt")
    ap.add_argument("--json", metavar="OUT.json", help="write structured findings as JSON")
    ap.add_argument("--apply", metavar="OUT.docx", help="write a copy with AUTO fixes applied")
    args = ap.parse_args(argv)

    src = Path(args.manuscript)
    if not src.exists():
        print(f"error: {src} not found", file=sys.stderr)
        return 2

    paragraphs = _read_paragraphs(src)
    findings = scan(paragraphs)

    if args.json:
        Path(args.json).write_text(
            json.dumps({"source": str(src),
                        "counts": {"auto": sum(f.mode == "auto" for f in findings),
                                   "review": sum(f.mode == "review" for f in findings),
                                   "total": len(findings)},
                        "findings": [asdict(f) for f in findings]}, indent=2),
            encoding="utf-8")
        print(f"wrote {args.json} ({len(findings)} findings)")

    if args.apply:
        if src.suffix.lower() == ".docx":
            n = _write_docx_auto(src, Path(args.apply))
        else:
            new_paras, n = apply_auto(paragraphs)
            Path(args.apply).write_text("\n".join(new_paras), encoding="utf-8")
        print(f"applied {n} auto fixes → {args.apply}")

    if not args.json and not args.apply:
        print(report(findings))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
